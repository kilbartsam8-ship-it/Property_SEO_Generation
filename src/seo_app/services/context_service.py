import io
from pathlib import Path
from typing import Any, BinaryIO

import pandas as pd
import requests

from ..config import Settings
from ..extractors.pdf import extract_pdf_text, extract_text_from_path
from ..extractors.url import UrlExtractionError, extract_website_text
from .info_extractor import ExtractionError, extract_info_from_text
from .summarizer import summarize_pdf_text


class ContextService:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def extract_path_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text = extract_text_from_path(path)
            return self._summarize_if_needed(text)
        if suffix == ".docx":
            try:
                import docx  # type: ignore
            except Exception as exc:
                raise ValueError("DOCX support requires package 'python-docx'.") from exc
            doc = docx.Document(str(path))
            return "\n".join(paragraph.text for paragraph in doc.paragraphs)
        if suffix == ".csv":
            dataframe = pd.read_csv(path)
            return dataframe.to_string(index=False)
        raise ValueError(f"Unsupported file type: {suffix}")

    def extract_upload_text(self, filename: str, file_obj: BinaryIO) -> str:
        lower_name = (filename or "").lower()
        if lower_name.endswith(".pdf"):
            text = extract_pdf_text(file_obj)
            return self._summarize_if_needed(text)
        if lower_name.endswith(".docx"):
            try:
                import docx  # type: ignore
            except Exception as exc:
                raise ValueError("DOCX support requires package 'python-docx'.") from exc
            file_obj.seek(0)
            doc = docx.Document(file_obj)
            return "\n".join(paragraph.text for paragraph in doc.paragraphs)
        if lower_name.endswith(".csv"):
            file_obj.seek(0)
            dataframe = pd.read_csv(file_obj)
            return dataframe.to_string(index=False)
        raise ValueError("Unsupported file type. Allowed: .pdf, .docx, .csv")

    def extract_pdf_url_text(self, pdf_url: str, timeout_seconds: int = 20) -> str:
        cleaned_url = (pdf_url or "").strip()
        if not cleaned_url:
            return ""

        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
            )
        }

        try:
            response = requests.get(cleaned_url, timeout=timeout_seconds, headers=headers)
            response.raise_for_status()
            buffer = io.BytesIO(response.content)
            text = extract_pdf_text(buffer)
            return self._summarize_if_needed(text)
        except requests.RequestException as exc:
            raise ValueError(f"Unable to fetch PDF from URL: {exc}") from exc
        except Exception as exc:
            raise ValueError(f"Unable to extract text from PDF URL: {exc}") from exc

    def extract_structured_info(self, text: str) -> dict[str, Any]:
        if not (text or "").strip():
            return {}
        try:
            extracted = extract_info_from_text(text, self.settings.groq_api_key, model=self.settings.model)
            return extracted if isinstance(extracted, dict) else {}
        except ExtractionError:
            return {}

    @staticmethod
    def extract_website_context(url: str) -> str:
        if not (url or "").strip():
            return ""
        try:
            return extract_website_text(url)
        except UrlExtractionError:
            return ""

    def _summarize_if_needed(self, text: str) -> str:
        if self.settings.summarize_long_pdf and len(text) > self.settings.summary_threshold:
            return summarize_pdf_text(text, self.settings.groq_api_key)
        return text
